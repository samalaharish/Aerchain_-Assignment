from pathlib import Path

from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "fixtures" / "vendor-responses"


def make_excel() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Commercial Quote"
    ws.append(["RFx Line", "Supplier Item", "Qty", "Price", "Currency", "UOM", "Freight", "Lead Time", "Quality"])
    rows = [
        [1, "3-ply brown shipping carton 200x150x120 mm", 6000, 780, "INR", "per 100 pcs", "Included", 21, "PASS"],
        [2, "3-ply brown shipping carton 250x180x150 mm", 6350, 822, "INR", "per 100 pcs", "Included", 21, "PASS"],
        [3, "3-ply brown shipping carton 300x220x180 mm", 6700, 864, "INR", "per 100 pcs", "Included", 21, "PASS"],
        [4, "3-ply printed retail carton 180x120x90 mm", 7050, 906, "INR", "per 100 pcs", "Included", 21, "PASS"],
        [5, "3-ply printed retail carton 220x160x120 mm", 7400, 948, "INR", "per 100 pcs", "Included", 21, "PASS"],
        [6, "5-ply heavy duty carton 300x250x220 mm", 7750, 990, "INR", "per 100 pcs", "Included", 21, "PASS"],
        [7, "5-ply heavy duty carton 400x300x250 mm", 8100, 1032, "INR", "per 100 pcs", "Included", 21, "PASS"],
        [8, "5-ply export carton 500x400x350 mm", 8450, 1074, "INR", "per 100 pcs", "Included", 21, "PASS"],
    ]
    for row in rows:
        ws.append(row)
    q = wb.create_sheet("Quality")
    q.append(["Question", "Answer", "Status"])
    q.append(["ISO 9001 certificate", "Valid through 2027", "PASS"])
    q.append(["Burst strength confirmation", "Attached test report AP-44ECT", "PASS"])
    wb.save(OUT / "vendor-a.xlsx")


def make_pdf() -> None:
    pdf = canvas.Canvas(str(OUT / "vendor-b.pdf"), pagesize=letter)
    pdf.setTitle("Bharat Corrugates Quote")
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(72, 740, "Bharat Corrugates - Corrugated Packaging Quotation")
    pdf.setFont("Helvetica", 9)
    pdf.drawString(72, 720, "RFx: rfx-corrugated-2026 | Currency INR unless noted")
    y = 690
    rows = [
        "Line 9 | C. board mailer small | 11.25 INR / piece | freight extra per footnote | 28 days",
        "Line 10 | C. board mailer medium | 11.90 INR / piece | freight 90 INR / 100 pcs | 28 days",
        "Line 11 | Die cut ecommerce shipper A | 12.55 INR / piece | freight extra per footnote | 28 days",
        "Line 12 | Die cut ecommerce shipper B | 13.20 INR / piece | freight 90 INR / 100 pcs | 28 days",
        "Line 13 | Partition insert six cell | 13.85 INR / piece | freight extra per footnote | lead time not stated",
        "Line 14 | Partition insert twelve cell | 14.50 INR / piece | freight 90 INR / 100 pcs | 28 days",
    ]
    for row in rows:
        pdf.drawString(72, y, row)
        y -= 22
    pdf.setFont("Helvetica-Oblique", 8)
    pdf.drawString(72, y - 10, "Footnote: freight extra for non-local delivery; exact freight subject to lane confirmation.")
    pdf.drawString(72, y - 24, "Quality: ISO certificate attached; ECT confirmation attached.")
    pdf.showPage()
    pdf.save()


def make_docx() -> None:
    doc = Document()
    doc.add_heading("CartonCraft Works - Email Follow-up", level=1)
    doc.add_paragraph("We can cover most sleeve and roll items. Commercials are below; freight to be discussed.")
    table = doc.add_table(rows=1, cols=6)
    headers = ["Line", "Description", "Qty", "Price", "Currency", "Lead time"]
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    rows = [
        [15, "Corrugated sleeve 1L bottle", 10900, 14.5, "INR", "35 days"],
        [16, "Corrugated sleeve 2L bottle", 11250, "", "INR", "35 days"],
        [17, "Edge protector 50 mm", "", 16.5, "INR", "35 days"],
        [18, "Edge protector 75 mm", 11950, 17.5, "", ""],
        [19, "Single face corrugated roll 36 inch", 12300, 18.5, "INR", ""],
        [20, "Single face corrugated roll 48 inch", 12650, 19.5, "INR", ""],
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    doc.add_paragraph("Quality questionnaire: ISO certificate renewal in progress; latest certificate not attached.")
    doc.add_paragraph("Note: line 20 quoted per box in internal sheet, pack size not confirmed.")
    doc.save(OUT / "vendor-c.docx")


def make_text() -> None:
    text = """Delta Fibreboard response
Currency: USD
Quality note: can meet most requested grades, but 44 ECT certificate is not current.

Top pad big size - maps maybe line 21 - 0.22 USD / piece - freight 0.01 USD / piece - 24 days
Separator pad 800 by 1200 - line 22 - 0.26 USD / piece - freight 0.01 USD / piece - 24 days
Pizza 9 plain - line 23 - price shown as 0 USD? please confirm - freight 0.01 USD / piece
Pizza 12 print - line 24 - 0.34 USD / unknown unit - freight 0.01 USD / piece
Fruit tray 2kg - line 25 - 0.38 USD / piece - duplicate desc fruit tray
Special fruit tray heavy - unmapped - 0.42 USD / piece
"""
    (OUT / "vendor-d.txt").write_text(text, encoding="utf-8")


def make_png() -> None:
    image = Image.new("RGB", (1100, 760), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default()
    lines = [
        "Eastern Box Makers - scanned rate card",
        "Display shpr sm - box INR 16 - pack unclear - lead time blank",
        "Display shipper large - INR 18 / box of 50 - freight missing - 32 days",
        "Mailer envelope rigid A4 - currency not visible - bundle price 20",
        "Returnable tote liner - INR 22 / bundle - freight included - 32 days",
        "Quality papers: not sent with this photo",
    ]
    y = 60
    for line in lines:
        draw.text((80, y), line, fill="black", font=font)
        y += 70
    image = image.rotate(-2, expand=True, fillcolor="white")
    image.save(OUT / "vendor-e.png")


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    make_excel()
    make_pdf()
    make_docx()
    make_text()
    make_png()


if __name__ == "__main__":
    main()
